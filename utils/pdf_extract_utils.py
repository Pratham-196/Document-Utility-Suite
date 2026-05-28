"""PDF page extraction utilities"""
from pypdf import PdfReader, PdfWriter
import os

def extract_pdf_pages(input_path, output_path, page_numbers):
    """
    Extract specific pages from PDF
    
    Args:
        input_path: Path to input PDF
        output_path: Path to save extracted pages
        page_numbers: List of page numbers to extract (1-indexed)
    
    Returns:
        output_path if successful, None otherwise
    """
    try:
        reader = PdfReader(input_path)
        writer = PdfWriter()
        
        total_pages = len(reader.pages)
        extracted_count = 0
        
        for page_num in page_numbers:
            if 1 <= page_num <= total_pages:
                writer.add_page(reader.pages[page_num - 1])
                extracted_count += 1
            else:
                print(f"Warning: Page {page_num} is out of range (1-{total_pages}), skipping...")
        
        if extracted_count == 0:
            print(f"Error: No valid pages to extract. Total pages in PDF: {total_pages}")
            return None
        
        with open(output_path, 'wb') as output_file:
            writer.write(output_file)
        
        print(f"Successfully extracted {extracted_count} page(s)")
        return output_path
    except Exception as e:
        print(f"Error extracting pages: {e}")
        import traceback
        traceback.print_exc()
        return None

def crop_pdf(input_path, output_path, left, top, right, bottom):
    """
    Crop all pages in PDF to specified dimensions
    
    Args:
        input_path: Path to input PDF
        output_path: Path to save cropped PDF
        left, top, right, bottom: Crop coordinates in points
    
    Returns:
        output_path if successful, None otherwise
    """
    try:
        import fitz
        
        doc = fitz.open(input_path)
        
        for page in doc:
            # Get current page dimensions
            rect = page.rect
            
            # Create crop box (left, top, right, bottom)
            crop_rect = fitz.Rect(left, top, right, bottom)
            
            # Ensure crop box is within page bounds
            crop_rect = crop_rect & rect
            
            # Set the crop box
            page.set_cropbox(crop_rect)
        
        doc.save(output_path)
        doc.close()
        return output_path
    except Exception as e:
        print(f"Error cropping PDF: {e}")
        return None

def pdf_to_ppt(input_path, output_path):
    """
    Convert PDF to PowerPoint (basic conversion)
    Note: This is a simplified version. Full conversion requires complex libraries.
    
    Args:
        input_path: Path to input PDF
        output_path: Path to save PPT file
    
    Returns:
        output_path if successful, None otherwise
    """
    try:
        from pptx import Presentation
        from pptx.util import Inches
        import fitz
        
        # Create presentation
        prs = Presentation()
        prs.slide_width = Inches(10)
        prs.slide_height = Inches(7.5)
        
        # Open PDF
        doc = fitz.open(input_path)
        
        # Convert each page to image and add to slide
        upload_dir = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'static', 'uploads')
        os.makedirs(upload_dir, exist_ok=True)

        for page_num in range(len(doc)):
            page = doc.load_page(page_num)
            
            # Render page as image
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
            img_path = os.path.join(upload_dir, f"temp_ppt_page_{page_num}.png")
            pix.save(img_path)
            
            # Add blank slide
            blank_slide_layout = prs.slide_layouts[6]  # Blank layout
            slide = prs.slides.add_slide(blank_slide_layout)
            
            # Add image to slide
            left = Inches(0)
            top = Inches(0)
            slide.shapes.add_picture(img_path, left, top, width=prs.slide_width)
            
            # Clean up temp image
            os.remove(img_path)
        
        doc.close()
        
        # Save presentation
        prs.save(output_path)
        return output_path
    except Exception as e:
        print(f"Error converting PDF to PPT: {e}")
        return None


def pdf_to_text(input_path, output_path):
    """Extract text layer from PDF directly (no OCR)."""
    try:
        import fitz
        doc = fitz.open(input_path)
        lines = []
        for i, page in enumerate(doc):
            lines.append(f"--- Page {i + 1} ---\n")
            lines.append(page.get_text())
            lines.append("\n")
        doc.close()
        text = "\n".join(lines)
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(text)
        return output_path
    except Exception as e:
        print(f"Error extracting text: {e}")
        return None


def pdf_to_csv(input_path, output_dir):
    """Extract tables from PDF pages to CSV files using pdfplumber."""
    try:
        import pdfplumber
        import csv

        os.makedirs(output_dir, exist_ok=True)
        output_files = []

        with pdfplumber.open(input_path) as pdf:
            for i, page in enumerate(pdf.pages):
                tables = page.extract_tables()
                for j, table in enumerate(tables):
                    csv_path = os.path.join(output_dir, f"table_p{i+1}_t{j+1}.csv")
                    with open(csv_path, 'w', newline='', encoding='utf-8') as f:
                        writer = csv.writer(f)
                        for row in table:
                            writer.writerow([cell or '' for cell in row])
                    output_files.append(csv_path)

        return output_files if output_files else None
    except ImportError:
        print("pdfplumber not installed. Run: pip install pdfplumber")
        return None
    except Exception as e:
        print(f"Error extracting tables: {e}")
        return None
